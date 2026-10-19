"""Expert (coach) review and approval form for the app's feedback, the live-test protocol and the usability
questionnaire (thesis Table 3.4, Section 3.6) -> Coach_Expert_Review_Form.docx

Usage:  .venv/Scripts/python scripts/make_expert_form.py
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import ROOT

OUT = ROOT / "forms" / "Coach_Expert_Review_Form.docx"
NAVY = RGBColor(0x1F, 0x3B, 0x6E)
BOX = "☐"
DECISION = f"{BOX} Approve\n{BOX} Approve with changes\n{BOX} Not approved"


def shade(cell, color="E8EDF5"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def para(doc, text="", bold=False, italic=False, size=10, after=4, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def heading(doc, text):
    p = para(doc, text, bold=True, size=12, after=4, color=NAVY)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True


def table(doc, head, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(size)
        shade(c)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v)
            r.font.size = Pt(size)
    for row in t.rows:
        tr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        tr.append(cs)
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    rep = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(rep)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)

    para(doc, "Expert Review and Approval Form", bold=True, size=16, after=2, color=NAVY)
    para(doc, "Landing Form Coach: posture feedback, live-test protocol and usability questionnaire",
         size=11, after=8, color=RGBColor(0x59, 0x59, 0x59))

    para(doc, "Research project: Detecting Wrong Posture During Plyometric Exercises")
    para(doc, "Student: H.A.D.P.C. Hettiarachchi (COBSCCOMP242P-033), BSc (Hons) Computing, NIBM / Coventry University")
    para(doc, "Supervisor: Dr. Thisara Weerasingha", after=8)

    para(doc, "Dear Coach,", after=6)
    para(doc, "Thank you for agreeing to review my final-year research project. I have built a web application that uses "
              "a laptop or phone camera to watch a person land from a box jump or depth jump and tells them whether the "
              "landing showed knee valgus (knees caving inward) or trunk lean (the trunk tilting to one side). The "
              "application runs entirely in the browser and no video is stored or sent anywhere.", after=6)
    para(doc, "Before I test the application with participants, I would like an experienced coach to check that "
              "(1) the faults are defined the way a coach would judge them, (2) the feedback messages are correct and "
              "safe, (3) the live-test plan is safe, and (4) the questions in my usability questionnaire are relevant "
              "and clear. Please mark one box for each item and add any comments or suggested wording. Your honest "
              "opinion is what matters: \"Not approved\" or a suggested change is just as useful as an approval. "
              "The review takes about 15 to 20 minutes.", after=6)
    para(doc, "Your name will not appear in the thesis unless you agree below; only your role and experience will be "
              "reported.", italic=True, after=8)

    heading(doc, "Reviewer details")
    table(doc, ["Detail", "Your answer"], [
        ["Full name", ""], ["Role / position", ""], ["Organisation or club", ""],
        ["Years of coaching experience", ""], ["Qualifications or certifications", ""],
        ["Experience with plyometric or landing training", ""],
        ["Date of review", ""],
    ], [6.0, 11.0], size=10)

    heading(doc, "Part 1: Fault definitions")
    para(doc, "The application judges each landing as one of three results. Do these definitions match how you would "
              "judge a landing when watching from the front?", after=4)
    table(doc, ["#", "Item", "Your decision", "Comments / suggested change"], [
        ["1.1", "Correct form: both knees stay in line over the toes and the trunk stays upright over the hips.",
         DECISION, ""],
        ["1.2", "Knee valgus: one or both knees move inward, past the line of the toes, when seen from the front.",
         DECISION, ""],
        ["1.3", "Trunk lean: the chest/trunk tilts to one side, away from upright over the hips, when seen from the "
                "front.", DECISION, ""],
        ["1.4", "The landing is judged from the first foot contact with the ground to the deepest point of the knee "
                "bend (maximum knee flexion).", DECISION, ""],
        ["1.5", "If a landing shows both faults, the more obvious one is reported.", DECISION, ""],
    ], [0.9, 7.0, 3.8, 5.3])

    heading(doc, "Part 2: Feedback messages")
    para(doc, "After each landing the application shows one of these messages on screen and can speak a short cue "
              "aloud. Are they correct, clear and what a coach would say?", after=4)
    table(doc, ["#", "Result", "On-screen message", "Spoken cue", "Your decision", "Comments / suggested wording"], [
        ["2.1", "Good landing", "Knees over toes, trunk upright. Keep it up!", "\"Good landing\"", DECISION, ""],
        ["2.2", "Knees caving in", "Push your knees out so they stay over your toes as you land.", "\"Knees out\"",
         DECISION, ""],
        ["2.3", "Leaning to the side", "Keep your chest centred over your hips when you land.", "\"Chest centred\"",
         DECISION, ""],
        ["2.4", "Timing", "The message appears about half a second after landing, before the next jump.", "–",
         DECISION, ""],
    ], [0.9, 2.3, 4.2, 2.2, 3.2, 4.2])

    heading(doc, "Part 3: Live-test plan and safety")
    para(doc, "Each of about 15 adult participants (18 or older) will complete 10 landings in front of the camera. "
              "A coach who cannot see the screen will rate every landing, and the ratings will be compared with the "
              "application's results.", after=4)
    table(doc, ["#", "Item", "Your decision", "Comments / suggested change"], [
        ["3.1", "Screening: a PAR-Q health questionnaire before starting; anyone with a current lower-limb injury or a "
                "history of ACL injury is excluded.", DECISION, ""],
        ["3.2", "A warm-up before the jumps; the test stops immediately if the participant reports any pain.",
         DECISION, ""],
        ["3.3", "6 landings with natural technique: 3 box jumps and 3 depth jumps from a 30 cm box.", DECISION, ""],
        ["3.4", "2 landings with a mild knee valgus and 2 with a mild trunk lean, performed on purpose after a "
                "demonstration; participants may decline these.", DECISION, ""],
        ["3.5", "Camera about 3 m in front of the participant at hip height (frontal view).", DECISION, ""],
    ], [0.9, 7.0, 3.8, 5.3])

    heading(doc, "Part 4: Usability questionnaire (completed by participants after using the application)")
    para(doc, "Please rate how relevant each question is to evaluating this application (1 = not relevant, "
              "2 = somewhat relevant, 3 = quite relevant, 4 = highly relevant) and whether it is clear to an athlete "
              "or coach. Suggest new wording for any unclear question.", after=4)
    R = f"{BOX}1 {BOX}2 {BOX}3 {BOX}4"
    C = f"{BOX}\u00a0Yes\u00a0\u00a0{BOX}\u00a0No"
    q = [
        ("A1", "What is your role? (Athlete / Coach or trainer / Both)"),
        ("A2", "What is your age group?"),
        ("A3", "How many years of experience do you have with plyometric (jump) training?"),
        ("A4", "How often do you use fitness or health-tracking apps?"),
        ("A5", "Which device did you use to test the application?"),
        ("C1", "The application correctly identified when my knees caved inward on landing."),
        ("C2", "The application correctly identified when my trunk leaned too far."),
        ("C3", "The feedback appeared quickly enough to be useful during the set."),
        ("C4", "The correction messages were clear and easy to understand."),
        ("C5", "The application worked well for checking my posture during plyometric exercises."),
        ("C6", "The application was reliable."),
        ("C7", "Using the application would improve my performance in plyometric training."),
        ("C8", "I find the application useful for my plyometric training."),
        ("C9", "I intend to continue using the application."),
        ("C10", "I predict that I would continue using the application."),
        ("C11", "Did the application give any result that you believe was wrong? (Yes / No / Not sure)"),
        ("C12", "If yes, what kind of wrong result did you see? (false knee valgus, false trunk lean, missed knee "
                "valgus, missed trunk lean, feedback too late, other)"),
        ("D1", "What did you like most about the application?"),
        ("D2", "Which part of the application was hardest to use?"),
        ("D3", "What one feature should be added or improved?"),
    ]
    table(doc, ["Item", "Question", "Relevance", "Clear?", "Suggested wording / comment"],
          [[k, v, R, C, ""] for k, v in q], [1.0, 6.5, 2.8, 2.4, 4.3])
    para(doc, "Section B (items B1 to B10) is the System Usability Scale (Brooke, 1996), a standard published "
              "questionnaire. Its wording is kept unchanged so that the score can be compared with other studies, so "
              "it does not need to be rated. Any general comment on it is welcome below.", italic=True, size=9, after=6)
    para(doc, "Is anything important missing from the questionnaire?", bold=True, after=2)
    for _ in range(3):
        para(doc, "_" * 98, after=2, color=RGBColor(0xA6, 0xA6, 0xA6))

    heading(doc, "Overall decision and declaration")
    para(doc, f"{BOX} I approve the fault definitions, feedback messages, live-test plan and questionnaire as they are.",
         after=3)
    para(doc, f"{BOX} I approve them, subject to the changes written in my comments above.", after=3)
    para(doc, f"{BOX} I do not approve them in their current form (reasons given above).", after=8)
    para(doc, "I confirm that I reviewed the items above independently, based on my own coaching experience, and that "
              "the ratings and comments on this form are my own.", after=6)
    para(doc, f"{BOX} I agree that my name may be listed in the thesis acknowledgements.   "
              f"{BOX} Please report only my role and experience.", after=14)
    sig = doc.add_table(rows=2, cols=2)
    for i, (a, b) in enumerate([("Reviewer signature", "Date"), ("Student signature", "Date")]):
        for j, lab in enumerate((a, b)):
            c = sig.rows[i].cells[j]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(18)
            p.add_run(("_" * 42 if j == 0 else "_" * 22) + "\n").font.size = Pt(10)
            r = p.add_run(lab)
            r.font.size = Pt(9)
            c.width = Cm(11.0 if j == 0 else 6.0)

    for s in doc.sections:
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Expert review form – Detecting Wrong Posture During Plyometric Exercises – H.A.D.P.C. "
                       "Hettiarachchi")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x80, 0x80, 0x80)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
