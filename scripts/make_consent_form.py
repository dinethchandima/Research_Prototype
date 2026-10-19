"""Two printable forms for the live test (thesis Section 3.12), one copy per participant:
  Live_Test_Consent_Form.docx   information sheet + consent (jump test only; the usability questionnaire is a
                                Google Form with its own consent question)
  PAR-Q_Form.docx               health screening before jumping

Usage:  .venv/Scripts/python scripts/make_consent_form.py
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

from common import ROOT
from make_expert_form import BOX, NAVY, heading, para, table

GREY = RGBColor(0x59, 0x59, 0x59)
PROJECT = "Detecting Wrong Posture During Plyometric Exercises"


def bullet(doc, text):
    p = para(doc, "•  " + text, after=2)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.4)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def title(doc, text, sub):
    para(doc, text, bold=True, size=15, after=1, color=NAVY)
    para(doc, sub, size=10, after=8, color=GREY)


def new_doc(footer):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(2.0)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(f"{PROJECT} – {footer}")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x80, 0x80, 0x80)
    return doc


def main():
    project = PROJECT
    doc = new_doc("participant information and consent (live test)")

    # ---------------------------------------------------------------- 1. information sheet
    title(doc, "Participant Information Sheet", f"Research project: {project}")
    para(doc, "You are invited to take part in a final-year research project for the BSc (Hons) Computing degree at "
              "NIBM (awarded by Coventry University). Please read this sheet before deciding. You may keep it.",
         after=6)
    heading(doc, "What is the study about?")
    para(doc, "I have built a web application that watches a person land from a jump through a laptop camera and gives "
              "a short message if the landing shows knee valgus (knees caving inward) or trunk lean (the upper body "
              "tilting to one side). This study tests how well the application's results agree with a coach's "
              "judgement and how easy it is to use.")
    heading(doc, "What will I be asked to do? (about 15 to 20 minutes)")
    bullet(doc, "Complete a short health check (PAR-Q) to make sure the jumps are safe for you.")
    bullet(doc, "Warm up, then perform 10 single jumps in front of the camera, with a rest between each: 3 box jumps "
                "onto a 30 cm box and 3 depth jumps (stepping off the box and landing on the floor) with your normal "
                "technique, then 2 landings with your knees slightly caving inward and 2 with your upper body slightly "
                "leaning to one side. These 4 will be demonstrated first, and you may decline them.")
    bullet(doc, "A coach will watch and rate each landing. The coach will not see the application's screen.")
    bullet(doc, "Afterwards, you will be invited to fill in a short online questionnaire about the application. It "
                "has its own consent question and is optional.")
    heading(doc, "Are there any risks?")
    para(doc, "The jumps are low and similar to normal training, but, as with any exercise, there is a small risk of "
              "muscle strain, a twisted ankle or losing balance. The intentionally poor landings will be mild and "
              "few. You will warm up first, the area will be clear, and you can stop at any moment. Tell the "
              "researcher immediately if you feel any pain or discomfort; the test will stop.")
    heading(doc, "What happens to my data?")
    bullet(doc, "No video is recorded or stored. The camera image is processed only inside the web browser on the "
                "researcher's laptop and is never uploaded.")
    bullet(doc, "Only these results are kept: the coach's rating and the application's result for each landing, "
                "and timing values. They are linked to a participant number (for "
                "example P01), not to your name.")
    bullet(doc, "Your name appears only on this consent form, which is stored separately and securely by the "
                "researcher and destroyed after the project has been assessed.")
    bullet(doc, "Anonymous results may appear in the thesis. You will not be identifiable in it.")
    heading(doc, "Do I have to take part?")
    para(doc, "No. Taking part is voluntary. You must be 18 or older. You can stop at any time during the session "
              "without giving a reason, and your results will then be deleted. After the session your data cannot be "
              "removed, because it is stored only under a number and is no longer linked to you.")
    heading(doc, "Contact")
    para(doc, "Researcher: H.A.D.P.C. Hettiarachchi (COBSCCOMP242P-033), NIBM.   "
              "Contact: ______________________________", after=2)
    para(doc, "Supervisor: Dr. Thisara Weerasingha, NIBM.", after=2)

    # ---------------------------------------------------------------- 2. consent form
    page_break(doc)
    title(doc, "Participant Consent Form", f"Research project: {project}")
    para(doc, "Participant ID (filled in by the researcher):  P______", bold=True, after=8)
    para(doc, "Please tick each box if you agree.", after=4)
    stmts = [
        "I have read and understood the Participant Information Sheet and have had the chance to ask questions.",
        "I am 18 years of age or older.",
        "I understand that taking part is voluntary and that I can stop at any time during the session without "
        "giving a reason.",
        "I understand the risks of the jumps and I will tell the researcher immediately if I feel any pain.",
        "I agree to perform 10 jump landings in front of the camera, including up to 4 intentionally mild poor "
        "landings that I may decline.",
        "I agree that a coach will watch and rate my landings.",
        "I understand that no video of me will be recorded or stored, and that my results will be kept anonymously "
        "under a participant number.",
        "I agree that my anonymous results may be used in the thesis and in related academic work.",
    ]
    table(doc, ["", "Statement", f"Tick {BOX}"],
          [[str(i), s, BOX] for i, s in enumerate(stmts, 1)], [0.8, 14.2, 2.0], size=10)
    para(doc, "I agree to take part in this study.", bold=True, after=14)
    sig = doc.add_table(rows=2, cols=3)
    for i, (a, b, c) in enumerate([("Participant name", "Signature", "Date"),
                                   ("Researcher name", "Signature", "Date")]):
        for j, lab in enumerate((a, b, c)):
            cell = sig.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(20)
            p.add_run("_" * (30 if j < 2 else 18) + "\n").font.size = Pt(10)
            p.add_run(lab).font.size = Pt(9)
            cell.width = Cm([6.5, 6.5, 4.0][j])
    para(doc, "").paragraph_format.space_after = Pt(6)
    para(doc, "One copy for the participant (on request) and one for the researcher. Consent forms are stored "
              "separately from all study results.", italic=True, size=9, color=GREY)

    doc.save(ROOT / "forms" / "Live_Test_Consent_Form.docx")

    # ---------------------------------------------------------------- 3. PAR-Q
    doc = new_doc("PAR-Q health screening")
    title(doc, "Physical Activity Readiness Questionnaire (PAR-Q)",
          "Health screening before the jump test. Adapted from the PAR-Q (Canadian Society for Exercise "
          "Physiology), with two questions added for landing tasks (8 and 9).")
    para(doc, "Participant ID:  P______", bold=True, after=6)
    para(doc, "Please answer each question honestly with Yes or No.", after=4)
    yn = f"{BOX} Yes   {BOX} No"
    qs = [
        "Has a doctor ever said that you have a heart condition and that you should only do physical activity "
        "recommended by a doctor?",
        "Do you feel pain in your chest when you do physical activity?",
        "In the past month, have you had chest pain when you were not doing physical activity?",
        "Do you lose your balance because of dizziness, or do you ever lose consciousness?",
        "Do you have a bone or joint problem (for example back, hip, knee or ankle) that could be made worse by "
        "jumping?",
        "Is a doctor currently prescribing medicine for your blood pressure or a heart condition?",
        "Do you know of any other reason why you should not do physical activity today (for example pregnancy or a "
        "recent illness)?",
        "Do you currently have any injury or pain in your hips, knees, ankles or feet?",
        "Have you ever had an anterior cruciate ligament (ACL) injury, or surgery on your knee or ankle?",
    ]
    table(doc, ["#", "Question", "Answer"], [[str(i), q, yn] for i, q in enumerate(qs, 1)],
          [0.8, 13.2, 3.0], size=10)
    para(doc, "If you answered YES to any question, you will not take part in the jump test today. Thank you "
              "for your honesty.", bold=True, after=6)
    para(doc, "I confirm that my answers are true to the best of my knowledge.", after=14)
    sig = doc.add_table(rows=1, cols=2)
    for j, lab in enumerate(("Participant signature", "Date")):
        cell = sig.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.add_run("_" * (40 if j == 0 else 20) + "\n").font.size = Pt(10)
        p.add_run(lab).font.size = Pt(9)
        cell.width = Cm([11.0, 6.0][j])
    para(doc, "").paragraph_format.space_after = Pt(6)
    para(doc, "Researcher use only:   PAR-Q passed   " + yn + "      (record in Live_Test_Sheet.xlsx, "
              "Participants tab)", size=9, color=GREY)

    doc.save(ROOT / "forms" / "PAR-Q_Form.docx")
    print("saved Live_Test_Consent_Form.docx and PAR-Q_Form.docx")


if __name__ == "__main__":
    main()
